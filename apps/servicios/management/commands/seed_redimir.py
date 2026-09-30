import os
from decimal import Decimal
from datetime import datetime, date
import openpyxl
from django.core.management.base import BaseCommand
from django.conf import settings
from django.core.files.base import ContentFile
from apps.empresas.models import Empresa
from apps.servicios.models import CatalogoMaterialRetiro, TicketRetiro, DetalleMaterialTicket


class Command(BaseCommand):
    help = "Carga el catalogo de materiales e importa los tickets de pesaje con desgloses y fotos desde los Excel"

    def handle(self, *args, **options):
        self.stdout.write(self.style.NOTICE("=== Iniciando Seed Modulo REDIMIR: Retiros y Pesajes ==="))

        # 1. Poblar Catalogo de Materiales
        materiales_data = [
            ("Cartón", "reciclable", "kg", 1),
            ("Papel", "reciclable", "kg", 2),
            ("Botellas Plásticas / PET", "reciclable", "kg", 3),
            ("Film", "reciclable", "kg", 4),
            ("Latas de Aluminio", "reciclable", "kg", 5),
            ("Zuncho", "reciclable", "kg", 6),
            ("Pallets", "valorizable", "un", 7),
            ("Carretes de Madera", "valorizable", "un", 8),
            ("Tambores", "valorizable", "un", 9),
            ("PP (Polipropileno)", "reciclable", "kg", 10),
            ("Barras de Perforación", "valorizable", "kg", 11),
            ("Descarga y Acopio", "general", "kg", 12),
            ("Chatarra / Metales", "valorizable", "kg", 13),
            ("Residuos Domésticos / Asimilables", "general", "kg", 14),
            ("Residuos Generales", "general", "kg", 15),
        ]

        materiales_map = {}
        for nombre, cat, um, orden in materiales_data:
            mat, _ = CatalogoMaterialRetiro.objects.get_or_create(
                nombre=nombre,
                defaults={
                    "categoria": cat,
                    "unidad_medida": um,
                    "orden": orden,
                    "activo": True
                }
            )
            materiales_map[nombre] = mat
        self.stdout.write(self.style.SUCCESS(f"[OK] Catalogo de materiales listo ({len(materiales_map)} items)"))

        # 2. Asegurar Empresas Generadoras
        empresa_open = Empresa.objects.filter(nombre__icontains="open").first()
        if not empresa_open:
            empresa_open = Empresa.objects.create(
                nombre="OPEN PLAZA CALAMA",
                rut="76.999.888-1",
                email_contacto="administracion@openplaza.cl",
                rubro="retail",
                estado="aprobada",
                activa=True
            )
            self.stdout.write(self.style.SUCCESS("[OK] Creada empresa OPEN PLAZA CALAMA"))
        else:
            self.stdout.write(self.style.SUCCESS(f"[OK] Empresa Open encontrada: {empresa_open.nombre}"))

        empresa_enaex = Empresa.objects.filter(nombre__icontains="enaex").first()
        if not empresa_enaex:
            empresa_enaex = Empresa.objects.create(
                nombre="ENAEX SERVICIOS S.A.",
                rut="96.536.000-K",
                email_contacto="operaciones@enaex.com",
                rubro="mineria",
                estado="aprobada",
                activa=True
            )
            self.stdout.write(self.style.SUCCESS("[OK] Creada empresa ENAEX SERVICIOS S.A."))
        else:
            self.stdout.write(self.style.SUCCESS(f"[OK] Empresa Enaex encontrada: {empresa_enaex.nombre}"))

        # Carpeta para respaldos
        media_respaldos = os.path.join(settings.MEDIA_ROOT, "tickets_retiro", "respaldos")
        os.makedirs(media_respaldos, exist_ok=True)

        base_dir = settings.BASE_DIR

        # 3. Importar Excel OPEN
        open_excel = os.path.join(base_dir, "Registro Tickets de Basura OPEN.xlsx")
        if os.path.exists(open_excel):
            self.import_open(open_excel, empresa_open, materiales_map, media_respaldos)
        else:
            self.stdout.write(self.style.WARNING(f"Archivo no encontrado: {open_excel}"))

        # 4. Importar Excel ENAEX Servicios Actualizado 2026
        enaex_actualizado = os.path.join(base_dir, "Registro servicios enaex actualizado 2026.xlsx")
        if os.path.exists(enaex_actualizado):
            self.import_enaex_actualizado(enaex_actualizado, empresa_enaex, materiales_map, media_respaldos)

        # 5. Importar Excel ENAEX Basura Mensual (con fotos)
        enaex_basura = os.path.join(base_dir, "Registro Tickets de Basura Enaex.xlsx")
        if os.path.exists(enaex_basura):
            self.import_enaex_basura_mensual(enaex_basura, empresa_enaex, materiales_map, media_respaldos)

        total_tickets = TicketRetiro.objects.count()
        total_detalles = DetalleMaterialTicket.objects.count()
        self.stdout.write(self.style.SUCCESS("=== SEED COMPLETADO CON EXITO ==="))
        self.stdout.write(self.style.SUCCESS(f"Total Tickets en BD: {total_tickets} | Total Desgloses: {total_detalles}"))

    def parse_float(self, val):
        if val is None or val == "":
            return Decimal("0")
        if isinstance(val, (int, float)):
            return Decimal(str(val))
        try:
            s = str(val).replace(".", "").replace(",", ".").strip()
            return Decimal(s)
        except Exception:
            return Decimal("0")

    def parse_date(self, val):
        if isinstance(val, (datetime, date)):
            return val.date() if isinstance(val, datetime) else val
        if isinstance(val, str):
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(val.strip(), fmt).date()
                except ValueError:
                    continue
        return date.today()

    def import_open(self, file_path, empresa, materiales_map, media_respaldos):
        self.stdout.write(f"Importando Open Calama desde {os.path.basename(file_path)}...")
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        images_by_row = {}
        for img in getattr(ws, "_images", []):
            try:
                row_idx = img.anchor._from.row + 1
                images_by_row[row_idx] = img
            except Exception:
                pass

        header = [str(cell.value).strip().upper() if cell.value is not None else "" for cell in ws[1]]
        col_map = {name: idx for idx, name in enumerate(header)}

        count = 0
        for row_idx in range(2, ws.max_row + 1):
            row = [cell.value for cell in ws[row_idx]]
            if not any(row):
                continue

            num_ticket = row[col_map.get("TICKET", 1)]
            if not num_ticket or str(num_ticket).strip() == "" or str(num_ticket).strip() == "None":
                continue
            num_ticket = str(num_ticket).strip()

            fecha = self.parse_date(row[col_map.get("FECHA", 0)])
            servicio = str(row[col_map.get("SERVICIO", 2)] or "Retiro de Residuos").strip()
            residuo = str(row[col_map.get("RESIDUO", 3)] or "Reciclables").strip()
            peso_ticket = self.parse_float(row[col_map.get("KILOS", 4)])

            ticket, created = TicketRetiro.objects.get_or_create(
                empresa=empresa,
                numero_ticket=num_ticket,
                defaults={
                    "fecha": fecha,
                    "faena_area": "Plaza Calama",
                    "tipo_servicio": servicio,
                    "tipo_residuo": residuo,
                    "peso_total_ticket": peso_ticket,
                }
            )
            if not created:
                ticket.fecha = fecha
                ticket.peso_total_ticket = peso_ticket
                ticket.save()

            if row_idx in images_by_row and not ticket.respaldo_ticket:
                try:
                    img = images_by_row[row_idx]
                    img_data = img._data()
                    ext = "png" if getattr(img, "format", "jpg").lower() == "png" else "jpg"
                    fname = f"ticket_open_{num_ticket}.{ext}"
                    ticket.respaldo_ticket.save(fname, ContentFile(img_data), save=True)
                except Exception as e:
                    self.stdout.write(self.style.WARNING(f"Foto ticket Open {num_ticket}: {e}"))

            mat_cols = [
                ("CARTON", "Cartón", "kg"),
                ("PAPEL", "Papel", "kg"),
                ("BOTELLAS", "Botellas Plásticas / PET", "kg"),
                ("FILM", "Film", "kg"),
                ("ALUMINIO", "Latas de Aluminio", "kg"),
                ("PALLET", "Pallets", "un"),
                ("CARRETES", "Carretes de Madera", "un"),
                ("ZUNCHO", "Zuncho", "kg"),
                ("TAMBORES", "Tambores", "un"),
            ]

            for col_key, mat_name, tipo in mat_cols:
                if col_key in col_map:
                    val = self.parse_float(row[col_map[col_key]])
                    if val > 0 and mat_name in materiales_map:
                        mat_obj = materiales_map[mat_name]
                        det, _ = DetalleMaterialTicket.objects.get_or_create(
                            ticket=ticket,
                            material=mat_obj,
                            defaults={
                                "peso_kg": val if tipo == "kg" else None,
                                "cantidad_unidades": int(val) if tipo == "un" else None,
                            }
                        )
                        if tipo == "kg":
                            det.peso_kg = val
                        else:
                            det.cantidad_unidades = int(val)
                        det.save()

            count += 1

        self.stdout.write(self.style.SUCCESS(f"[OK] Open Calama: {count} tickets procesados."))

    def import_enaex_actualizado(self, file_path, empresa, materiales_map, media_respaldos):
        self.stdout.write(f"Importando Enaex Actualizado desde {os.path.basename(file_path)}...")
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheets = wb.sheetnames

        total_enaex = 0
        for sheet_name in sheets:
            ws = wb[sheet_name]
            images_by_row = {}
            for img in getattr(ws, "_images", []):
                try:
                    row_idx = img.anchor._from.row + 1
                    images_by_row[row_idx] = img
                except Exception:
                    pass

            header_row = 1
            for r in range(1, 4):
                vals = [str(cell.value).upper() for cell in ws[r] if cell.value is not None]
                if any("TICKET" in v or "FOLIO" in v for v in vals):
                    header_row = r
                    break

            header = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[header_row]]
            h_lower = [h.lower() for h in header]

            col_ticket = next((i for i, h in enumerate(h_lower) if 'ticket' in h or 'folio' in h or h == ''), 0)
            col_fecha = next((i for i, h in enumerate(h_lower) if 'fecha' in h), 1)
            col_servicio = next((i for i, h in enumerate(h_lower) if 'servicio' in h), 2)
            col_residuo = next((i for i, h in enumerate(h_lower) if 'residuo' in h), 3)
            col_peso = next((i for i, h in enumerate(h_lower) if 'peso total' in h or 'total (kg)' in h), 4)

            # Mapeo exhaustivo de materiales
            mat_mappings = []
            def _find_mat(m_name):
                return materiales_map.get(m_name) or CatalogoMaterialRetiro.objects.filter(nombre__iexact=m_name).first()

            # Columnas directas de peso
            for key_term, m_name in [
                ('cart', 'Cartón'), ('papel', 'Papel'), ('pet', 'Botellas Plásticas / PET'),
                ('botella', 'Botellas Plásticas / PET'), ('film', 'Film'), ('aluminio', 'Latas de Aluminio'),
                ('lata', 'Latas de Aluminio'), ('suncho', 'Zuncho'), ('zuncho', 'Zuncho'),
                ('chatarra', 'Chatarra / Metales'), ('vidrio', 'Vidrio'), ('pp', 'PP (Polipropileno)'),
                ('barra', 'Barras de Perforación'), ('descarga', 'Descarga y Acopio')
            ]:
                idx = next((i for i, h in enumerate(h_lower) if key_term in h and 'cant' not in h), None)
                if idx is not None:
                    m_obj = _find_mat(m_name)
                    if m_obj:
                        mat_mappings.append(('peso', idx, None, m_obj))

            # Columnas combo (peso y/o cantidad con conversión)
            combos = [
                ('Pallets', 'peso palet', 'cantidad palet', 'mal'),
                ('Pallets en Mal Estado', 'peso palet', 'cantidad palet', None),
                ('Pallets Plásticos', 'pallet pl', 'pallet pl', None),
                ('Carretes de Madera', 'carrete.*madera', 'carrete.*madera', None),
                ('Carretes Plásticos', 'carrete.*pl', 'carrete.*pl', 'apd'),
                ('Carretes Plásticos APD', 'carrete.*apd', 'carrete.*apd', None),
                ('Tambores', 'peso.*tambor', 'tambores', '50'),
                ('Tambores 50 Litros', '50.*litro', '50.*litro', None),
                ('Tambor Metálico', 'metal', 'metal', None),
                ('Bidones 20 Litros', '20.*litro', '20.*litro', None),
                ('Bidones 25 Litros', '25.*litro', '25.*litro', None),
            ]
            for m_name, p_term, c_term, neg in combos:
                m_obj = _find_mat(m_name)
                if not m_obj: continue
                p_idx = next((i for i, h in enumerate(h_lower) if re.search(p_term, h) and 'peso' in h and (not neg or neg not in h)), None)
                c_idx = next((i for i, h in enumerate(h_lower) if (re.search(c_term, h) or h == c_term) and 'peso' not in h and (not neg or neg not in h)), None)
                if p_idx is not None or c_idx is not None:
                    mat_mappings.append(('combo', p_idx, c_idx, m_obj))

            for row_idx in range(header_row + 1, ws.max_row + 1):
                num_t = ws.cell(row=row_idx, column=col_ticket + 1).value
                if not num_t or str(num_t).strip() == "" or str(num_t).strip() == "None":
                    continue
                try:
                    num_ticket = str(int(float(num_t)))
                except Exception:
                    num_ticket = str(num_t).strip()

                fecha = self.parse_date(ws.cell(row=row_idx, column=col_fecha + 1).value)
                faena = str(ws.cell(row=row_idx, column=col_servicio + 1).value or "Faena Enaex").strip()
                residuo = str(ws.cell(row=row_idx, column=col_residuo + 1).value or "Reciclaje").strip()
                peso_ticket = self.parse_float(ws.cell(row=row_idx, column=col_peso + 1).value)

                ticket, created = TicketRetiro.objects.get_or_create(
                    empresa=empresa,
                    numero_ticket=num_ticket,
                    defaults={
                        "fecha": fecha,
                        "faena_area": "Faena Enaex",
                        "tipo_servicio": faena,
                        "tipo_residuo": residuo,
                        "peso_total_ticket": peso_ticket,
                    }
                )
                if not created:
                    ticket.fecha = fecha
                    ticket.tipo_servicio = faena
                    ticket.tipo_residuo = residuo
                    if peso_ticket > 0:
                        ticket.peso_total_ticket = peso_ticket
                    ticket.save()

                if row_idx in images_by_row and not ticket.respaldo_ticket:
                    try:
                        img = images_by_row[row_idx]
                        img_data = img._data()
                        ext = "png" if getattr(img, "format", "jpg").lower() == "png" else "jpg"
                        fname = f"ticket_enaex_{num_ticket}.{ext}"
                        ticket.respaldo_ticket.save(fname, ContentFile(img_data), save=True)
                    except Exception as e:
                        pass

                for m_type, p_idx, c_idx, mat_obj in mat_mappings:
                    peso_val = Decimal('0.00')
                    cant_val = Decimal('0.00')
                    if m_type == 'peso' and p_idx is not None:
                        peso_val = self.parse_float(ws.cell(row=row_idx, column=p_idx + 1).value)
                    elif m_type == 'combo':
                        if p_idx is not None:
                            peso_val = self.parse_float(ws.cell(row=row_idx, column=p_idx + 1).value)
                        if c_idx is not None:
                            cant_val = self.parse_float(ws.cell(row=row_idx, column=c_idx + 1).value)
                        if peso_val == 0 and cant_val > 0 and mat_obj.peso_unitario_kg > 0:
                            peso_val = cant_val * mat_obj.peso_unitario_kg

                    if peso_val > 0 or cant_val > 0:
                        det, _ = DetalleMaterialTicket.objects.get_or_create(
                            ticket=ticket,
                            material=mat_obj,
                            defaults={"peso_kg": peso_val, "cantidad_unidades": cant_val if cant_val > 0 else None}
                        )
                        det.peso_kg = peso_val
                        if cant_val > 0:
                            det.cantidad_unidades = cant_val
                        det.save()

                total_enaex += 1

        self.stdout.write(self.style.SUCCESS(f"[OK] Enaex Actualizado: {total_enaex} tickets procesados con desglose completo y conversión."))

    def import_enaex_basura_mensual(self, file_path, empresa, materiales_map, media_respaldos):
        self.stdout.write(f"Importando Enaex Basura Mensual desde {os.path.basename(file_path)}...")
        wb = openpyxl.load_workbook(file_path, data_only=True)
        count_mensual = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            images_by_row = {}
            for img in getattr(ws, "_images", []):
                try:
                    row_idx = img.anchor._from.row + 1
                    images_by_row[row_idx] = img
                except Exception:
                    pass

            header_row = None
            col_ticket = None
            col_peso = None
            col_fecha = None
            col_obs = None

            for r in range(1, 6):
                row_vals = [str(c.value).strip().upper() if c.value is not None else "" for c in ws[r]]
                for idx, v in enumerate(row_vals):
                    if "TICKET" in v:
                        col_ticket = idx
                        header_row = r
                    elif "PESO" in v or "KILO" in v:
                        col_peso = idx
                    elif "FECHA" in v:
                        col_fecha = idx
                    elif "OBS" in v:
                        col_obs = idx

                if header_row and col_ticket is not None:
                    break

            if not header_row or col_ticket is None:
                continue

            for r in range(header_row + 1, ws.max_row + 1):
                row_cells = [cell.value for cell in ws[r]]
                if not any(row_cells):
                    continue

                if col_ticket >= len(row_cells):
                    continue
                num_ticket = row_cells[col_ticket]
                if not num_ticket or str(num_ticket).strip() == "" or str(num_ticket).strip() == "None":
                    continue
                num_ticket = str(num_ticket).strip()

                fecha = date.today()
                if col_fecha is not None and col_fecha < len(row_cells):
                    fecha = self.parse_date(row_cells[col_fecha])

                peso = Decimal("0")
                if col_peso is not None and col_peso < len(row_cells):
                    peso = self.parse_float(row_cells[col_peso])

                obs = ""
                if col_obs is not None and col_obs < len(row_cells):
                    obs = str(row_cells[col_obs] or "").strip()

                ticket, created = TicketRetiro.objects.get_or_create(
                    empresa=empresa,
                    numero_ticket=num_ticket,
                    defaults={
                        "fecha": fecha,
                        "faena_area": "Planta Río Loa",
                        "tipo_servicio": "Retiro Residuos Domésticos",
                        "tipo_residuo": "Residuos Domésticos / Asimilables",
                        "peso_total_ticket": peso,
                        "observaciones": obs,
                    }
                )
                if not created:
                    if not ticket.respaldo_ticket and r in images_by_row:
                        pass
                else:
                    count_mensual += 1

                # Extraer foto si está disponible
                if r in images_by_row and not ticket.respaldo_ticket:
                    try:
                        img = images_by_row[r]
                        img_data = img._data()
                        ext = "png" if getattr(img, "format", "jpg").lower() == "png" else "jpg"
                        fname = f"ticket_domestico_{num_ticket}.{ext}"
                        ticket.respaldo_ticket.save(fname, ContentFile(img_data), save=True)
                    except Exception as e:
                        pass

                # Desglose de residuo domestico
                if peso > 0 and "Residuos Domésticos / Asimilables" in materiales_map:
                    mat_dom = materiales_map["Residuos Domésticos / Asimilables"]
                    det, _ = DetalleMaterialTicket.objects.get_or_create(
                        ticket=ticket,
                        material=mat_dom,
                        defaults={"peso_kg": peso}
                    )
                    det.peso_kg = peso
                    det.save()

        self.stdout.write(self.style.SUCCESS(f"[OK] Enaex Basura Mensual: {count_mensual} nuevos tickets importados con fotos."))
