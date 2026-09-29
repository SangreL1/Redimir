import os
import re
import json
import datetime
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.db.models import Q, Sum
from django.core.paginator import Paginator
from django.utils import timezone

from .models import (
    TicketRetiro, DetalleMaterialTicket, CatalogoMaterialRetiro,
    FotoRegistroRSD, FotoRegistroEscombros, FotoRegistroReciclables,
)
from .forms import TicketRetiroForm, DetalleMaterialTicketForm
from apps.empresas.models import Empresa


def _es_admin(user):
    return getattr(user, 'rol', '') == 'admin' or user.is_staff or user.is_superuser


@login_required
def retiros_lista(request):
    """
    Bandeja de Tickets y Retiros de Residuos.
    Multi-empresa: Los administradores ven todas las empresas; usuarios de empresa ven la suya.
    """
    tickets = TicketRetiro.objects.select_related('empresa').prefetch_related('detalles__material').all().order_by('-fecha', '-id')

    # Restricción multi-empresa
    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        tickets = tickets.filter(empresa=request.user.empresa)

    # Filtros
    q = request.GET.get('q', '').strip()
    empresa_id = request.GET.get('empresa', '').strip()
    faena_f = request.GET.get('faena', '').strip()
    servicio_f = request.GET.get('tipo_servicio', '').strip()
    residuo_f = request.GET.get('tipo_residuo', '').strip()
    desde_f = request.GET.get('desde', '').strip() or request.GET.get('fecha_desde', '').strip()
    hasta_f = request.GET.get('hasta', '').strip() or request.GET.get('fecha_hasta', '').strip()

    if q:
        tickets = tickets.filter(
            Q(numero_ticket__icontains=q) |
            Q(observaciones__icontains=q) |
            Q(empresa__nombre__icontains=q) |
            Q(faena_area__icontains=q)
        )
    if empresa_id:
        tickets = tickets.filter(empresa_id=empresa_id)
    if faena_f:
        tickets = tickets.filter(faena_area__icontains=faena_f)
    if servicio_f:
        tickets = tickets.filter(tipo_servicio__icontains=servicio_f)
    if residuo_f:
        tickets = tickets.filter(tipo_residuo__icontains=residuo_f)
    if desde_f:
        tickets = tickets.filter(fecha__gte=desde_f)
    if hasta_f:
        tickets = tickets.filter(fecha__lte=hasta_f)

    # Listas para opciones de filtros
    empresa_seleccionada_obj = None
    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        empresas_list = [request.user.empresa]
        empresa_seleccionada_obj = request.user.empresa
    else:
        empresas_list = Empresa.objects.filter(activa=True).order_by('nombre')
        if empresa_id:
            empresa_seleccionada_obj = Empresa.objects.filter(id=empresa_id).first()

    # Si hay una empresa seleccionada, listar sus servicios predeterminados + los registrados
    if empresa_seleccionada_obj:
        servs_base = list(empresa_seleccionada_obj.get_servicios_predeterminados_list())
        servs_db = list(TicketRetiro.objects.filter(empresa=empresa_seleccionada_obj).values_list('tipo_servicio', flat=True).exclude(tipo_servicio__isnull=True).exclude(tipo_servicio='').distinct())
        servicios_list = sorted(list(set(servs_base + servs_db)))
    else:
        servicios_db = list(TicketRetiro.objects.values_list('tipo_servicio', flat=True).exclude(tipo_servicio__isnull=True).exclude(tipo_servicio='').distinct())
        servicios_list = sorted(list(set(servicios_db + Empresa.SERVICIOS_ESTANDAR)))

    # Mapa de servicios por empresa para actualización dinámica en JS
    empresas_servicios_map = {
        str(emp.id): emp.get_servicios_predeterminados_list()
        for emp in empresas_list
    }

    residuos_list = TicketRetiro.objects.values_list('tipo_residuo', flat=True).exclude(tipo_residuo__isnull=True).exclude(tipo_residuo='').distinct().order_by('tipo_residuo')

    total_tickets = tickets.count()
    total_kilos = tickets.aggregate(total=Sum('peso_total_ticket'))['total'] or Decimal('0.00')

    detalles_qs = DetalleMaterialTicket.objects.filter(ticket__in=tickets)
    total_desglose = detalles_qs.aggregate(total=Sum('peso_kg'))['total'] or Decimal('0.00')
    total_unidades = detalles_qs.aggregate(total=Sum('cantidad_unidades'))['total'] or Decimal('0.00')

    kilos_reciclaje = detalles_qs.filter(material__categoria__icontains='recicl').aggregate(total=Sum('peso_kg'))['total'] or Decimal('0.00')
    kilos_maderas = detalles_qs.filter(material__categoria__icontains='valori').aggregate(total=Sum('peso_kg'))['total'] or Decimal('0.00')
    kilos_domesticos = detalles_qs.filter(material__categoria__icontains='gene').aggregate(total=Sum('peso_kg'))['total'] or Decimal('0.00')

    tickets_con_respaldo = tickets.exclude(respaldo_ticket='').exclude(respaldo_ticket__isnull=True).count()
    pct_respaldo = round((tickets_con_respaldo / total_tickets * 100), 1) if total_tickets > 0 else 0

    kpis = {
        'total_tickets': total_tickets,
        'total_kilos_ticket': total_kilos,
        'total_kilos_desglose': total_desglose,
        'total_unidades': total_unidades,
    }

    paginator = Paginator(tickets, 25)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'tickets': page_obj,
        'kpis': kpis,
        'total_tickets': total_tickets,
        'total_kilos': total_kilos,
        'total_desglose': total_desglose,
        'total_unidades': total_unidades,
        'kilos_reciclaje': kilos_reciclaje,
        'kilos_maderas': kilos_maderas,
        'kilos_domesticos': kilos_domesticos,
        'pct_respaldo': pct_respaldo,
        'empresas': empresas_list,
        'empresas_list': empresas_list,
        'servicios_list': servicios_list,
        'residuos_list': residuos_list,
        'empresas_servicios_json': json.dumps(empresas_servicios_map),
        'filtros': {
            'q': q,
            'empresa': empresa_id,
            'faena': faena_f,
            'tipo_servicio': servicio_f,
            'tipo_residuo': residuo_f,
            'desde': desde_f,
            'hasta': hasta_f,
        }
    }
    return render(request, 'servicios/retiros_lista.html', context)


@login_required
def retiro_crear(request):
    """
    Ingreso de nuevo retiro con desglose dinámico de materiales.
    """
    materiales_catalogo = CatalogoMaterialRetiro.objects.filter(activo=True).order_by('orden', 'nombre')

    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        empresas_disponibles = [request.user.empresa]
        empresa_default = request.user.empresa
    else:
        empresas_disponibles = Empresa.objects.filter(activa=True).order_by('nombre')
        empresa_default = empresas_disponibles.first()

    if request.method == 'POST':
        form = TicketRetiroForm(request.POST, request.FILES)
        if form.is_valid():
            ticket = form.save(commit=False)
            if getattr(request.user, 'rol', '') == 'empresa':
                ticket.empresa = request.user.empresa
            ticket.usuario_registro = request.user
            ticket.save()

            # Guardar líneas dinámicas de materiales
            mat_ids = request.POST.getlist('material_id[]')
            pesos = request.POST.getlist('peso_kg[]')
            cantidades = request.POST.getlist('cantidad_unidades[]')
            observaciones_items = request.POST.getlist('obs_detalle[]')

            for i in range(len(mat_ids)):
                mat_id = mat_ids[i]
                if not mat_id:
                    continue
                try:
                    mat_obj = CatalogoMaterialRetiro.objects.get(pk=mat_id)
                except CatalogoMaterialRetiro.DoesNotExist:
                    continue

                peso_val = Decimal('0.00')
                if i < len(pesos) and pesos[i].strip():
                    try:
                        peso_val = Decimal(pesos[i].strip())
                    except Exception:
                        peso_val = Decimal('0.00')

                cant_val = None
                if i < len(cantidades) and cantidades[i].strip():
                    try:
                        cant_val = Decimal(cantidades[i].strip())
                    except Exception:
                        cant_val = None

                obs_val = observaciones_items[i].strip() if i < len(observaciones_items) else ''

                if (peso_val and peso_val > 0) or (cant_val and cant_val > 0):
                    DetalleMaterialTicket.objects.create(
                        ticket=ticket,
                        material=mat_obj,
                        peso_kg=peso_val,
                        cantidad_unidades=cant_val,
                        observaciones=obs_val
                    )

            messages.success(request, f"✅ Ticket N° {ticket.numero_ticket} registrado exitosamente.")
        else:
            messages.error(request, "⚠️ Por favor corrige los errores en el formulario.")
    else:
        servicios_default = empresa_default.get_servicios_predeterminados_list() if empresa_default else Empresa.SERVICIOS_ESTANDAR
        form = TicketRetiroForm(initial={
            'fecha': timezone.now().date(),
            'empresa': empresa_default,
            'tipo_residuo': 'Reciclaje',
            'tipo_servicio': servicios_default[0] if servicios_default else 'Áreas Productivas'
        })

    empresas_servicios_map = {
        str(emp.id): emp.get_servicios_predeterminados_list()
        for emp in empresas_disponibles
    }

    return render(request, 'servicios/retiro_form.html', {
        'form': form,
        'materiales_catalogo': materiales_catalogo,
        'empresas': empresas_disponibles,
        'empresas_disponibles': empresas_disponibles,
        'empresas_servicios_json': json.dumps(empresas_servicios_map),
        'modo_edicion': False,
        'detalles_existentes': [],
    })


@login_required
def retiro_editar(request, pk):
    """
    Edición de un ticket existente y actualización de sus materiales.
    """
    ticket = get_object_or_404(TicketRetiro, pk=pk)
    materiales_catalogo = CatalogoMaterialRetiro.objects.filter(activo=True).order_by('orden', 'nombre')

    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        if ticket.empresa != request.user.empresa:
            messages.error(request, "Sin permisos para editar retiros de otra empresa.")
            return redirect('retiros-lista')
        empresas_disponibles = [request.user.empresa]
    else:
        empresas_disponibles = Empresa.objects.filter(activa=True).order_by('nombre')

    if request.method == 'POST':
        form = TicketRetiroForm(request.POST, request.FILES, instance=ticket)
        if form.is_valid():
            ticket = form.save()

            # Reemplazar líneas de desglose
            ticket.detalles.all().delete()

            mat_ids = request.POST.getlist('material_id[]')
            pesos = request.POST.getlist('peso_kg[]')
            cantidades = request.POST.getlist('cantidad_unidades[]')
            observaciones_items = request.POST.getlist('obs_detalle[]')

            for i in range(len(mat_ids)):
                mat_id = mat_ids[i]
                if not mat_id:
                    continue
                try:
                    mat_obj = CatalogoMaterialRetiro.objects.get(pk=mat_id)
                except CatalogoMaterialRetiro.DoesNotExist:
                    continue

                peso_val = Decimal('0.00')
                if i < len(pesos) and pesos[i].strip():
                    try:
                        peso_val = Decimal(pesos[i].strip())
                    except Exception:
                        peso_val = Decimal('0.00')

                cant_val = None
                if i < len(cantidades) and cantidades[i].strip():
                    try:
                        cant_val = Decimal(cantidades[i].strip())
                    except Exception:
                        cant_val = None

                obs_val = observaciones_items[i].strip() if i < len(observaciones_items) else ''

                if (peso_val and peso_val > 0) or (cant_val and cant_val > 0):
                    DetalleMaterialTicket.objects.create(
                        ticket=ticket,
                        material=mat_obj,
                        peso_kg=peso_val,
                        cantidad_unidades=cant_val,
                        observaciones=obs_val
                    )

            messages.success(request, f"✅ Ticket N° {ticket.numero_ticket} actualizado correctamente.")
            return redirect('retiro-detalle', pk=ticket.pk)
        else:
            messages.error(request, "⚠️ Error al guardar las modificaciones.")
    else:
        form = TicketRetiroForm(instance=ticket)

    detalles_existentes = ticket.detalles.select_related('material').all()

    empresas_servicios_map = {
        str(emp.id): emp.get_servicios_predeterminados_list()
        for emp in empresas_disponibles
    }

    return render(request, 'servicios/retiro_form.html', {
        'form': form,
        'ticket': ticket,
        'detalles_existentes': detalles_existentes,
        'materiales_catalogo': materiales_catalogo,
        'empresas': empresas_disponibles,
        'empresas_disponibles': empresas_disponibles,
        'empresas_servicios_json': json.dumps(empresas_servicios_map),
        'modo_edicion': True,
    })


@login_required
def retiro_detalle(request, pk):
    """
    Ficha de detalle técnico del ticket con visor de comprobante y desglose.
    """
    ticket = get_object_or_404(
        TicketRetiro.objects.select_related('empresa', 'usuario_registro').prefetch_related('detalles__material'),
        pk=pk
    )

    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        if ticket.empresa != request.user.empresa:
            messages.error(request, "Acceso no permitido.")
            return redirect('retiros-lista')

    detalles = list(ticket.detalles.all())
    peso_total = float(ticket.peso_total_ticket or 1.0)
    for d in detalles:
        d.porcentaje = round((float(d.peso_kg) / peso_total * 100), 1) if peso_total > 0 else 0

    return render(request, 'servicios/retiro_detalle.html', {
        'ticket': ticket,
        'detalles': detalles,
    })


@login_required
def retiro_eliminar(request, pk):
    """
    Eliminación de ticket.
    """
    ticket = get_object_or_404(TicketRetiro, pk=pk)

    if not _es_admin(request.user):
        messages.error(request, "Solo administradores pueden eliminar tickets.")
        return redirect('retiros-lista')

    numero = ticket.numero_ticket
    if request.method == 'POST':
        ticket.delete()
        messages.success(request, f"🗑️ Ticket N° {numero} eliminado exitosamente.")
        return redirect('retiros-lista')

    return render(request, 'servicios/retiro_confirmar_eliminar.html', {'ticket': ticket})


@login_required
def retiros_informe(request):
    """
    Generador de Informes Matriz de Residuos por Empresa y Fecha.
    """
    empresa_id = request.GET.get('empresa', '').strip()
    faena = request.GET.get('faena', '').strip()
    tipo_servicio = request.GET.get('tipo_servicio', '').strip()
    tipo_residuo = request.GET.get('tipo_residuo', '').strip()
    desde = request.GET.get('desde', '').strip() or request.GET.get('fecha_desde', '').strip()
    hasta = request.GET.get('hasta', '').strip() or request.GET.get('fecha_hasta', '').strip()

    qs = TicketRetiro.objects.select_related('empresa').prefetch_related('detalles__material').all().order_by('fecha', 'id')

    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        qs = qs.filter(empresa=request.user.empresa)
        empresas_list = [request.user.empresa]
        empresa_actual = request.user.empresa
    else:
        empresas_list = Empresa.objects.filter(activa=True).order_by('nombre')
        if empresa_id:
            qs = qs.filter(empresa_id=empresa_id)
            empresa_actual = Empresa.objects.filter(id=empresa_id).first()
        else:
            empresa_actual = None

    if faena:
        qs = qs.filter(faena_area__icontains=faena)
    if tipo_servicio:
        qs = qs.filter(tipo_servicio__icontains=tipo_servicio)
    if tipo_residuo:
        qs = qs.filter(tipo_residuo__icontains=tipo_residuo)
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)

    tickets = list(qs)

    materiales_presentes_ids = DetalleMaterialTicket.objects.filter(ticket__in=qs).values_list('material_id', flat=True).distinct()
    if materiales_presentes_ids:
        materiales = list(CatalogoMaterialRetiro.objects.filter(id__in=materiales_presentes_ids).order_by('orden', 'nombre'))
    else:
        materiales = list(CatalogoMaterialRetiro.objects.filter(activo=True).order_by('orden', 'nombre'))

    filas = []
    totales_por_mat = {mat.id: Decimal('0.00') for mat in materiales}
    total_peso_bascula = Decimal('0.00')
    total_desglose_general = Decimal('0.00')
    total_unidades_general = Decimal('0.00')

    for t in tickets:
        total_peso_bascula += (t.peso_total_ticket or Decimal('0.00'))
        det_map = {d.material_id: d for d in t.detalles.all()}
        
        mat_cols = []
        total_desglose_ticket = Decimal('0.00')
        
        for mat in materiales:
            det = det_map.get(mat.id)
            if det:
                if mat.unidad_medida == 'un':
                    val = det.cantidad_unidades or Decimal('0.00')
                    total_unidades_general += val
                else:
                    val = det.peso_kg or Decimal('0.00')
                    total_desglose_general += val
                    total_desglose_ticket += val
                
                totales_por_mat[mat.id] += val
                mat_cols.append({
                    'material': mat,
                    'valor': val,
                    'unidad': mat.unidad_medida,
                })
            else:
                mat_cols.append({
                    'material': mat,
                    'valor': Decimal('0.00'),
                    'unidad': mat.unidad_medida,
                })

        filas.append({
            'ticket': t,
            'materiales_cols': mat_cols,
            'total_desglose_kg': total_desglose_ticket,
        })

    totales_fila = {
        'total_tickets': len(tickets),
        'peso_total_ticket': total_peso_bascula,
        'total_desglose_kg': total_desglose_general,
        'total_unidades': total_unidades_general,
        'materiales': [{'valor': totales_por_mat[mat.id]} for mat in materiales],
    }

    servicios_list = TicketRetiro.objects.values_list('tipo_servicio', flat=True).exclude(tipo_servicio__isnull=True).exclude(tipo_servicio='').distinct().order_by('tipo_servicio')

    context = {
        'filas': filas,
        'materiales': materiales,
        'totales_fila': totales_fila,
        'tickets_count': len(tickets),
        'columnas_materiales': materiales,
        'matriz_filas': filas,
        'total_peso_bascula': total_peso_bascula,
        'empresas': empresas_list,
        'empresas_list': empresas_list,
        'empresa_seleccionada': empresa_id,
        'empresa_actual': empresa_actual,
        'fecha_desde': desde,
        'fecha_hasta': hasta,
        'servicios_list': servicios_list,
        'filtros': {
            'empresa': empresa_id,
            'faena': faena,
            'tipo_servicio': tipo_servicio,
            'tipo_residuo': tipo_residuo,
            'desde': desde,
            'hasta': hasta,
        }
    }
    return render(request, 'servicios/retiros_informe.html', context)


@login_required
def exportar_retiros_excel(request):
    """
    Descarga del Informe Consolidado en Excel (.xlsx) con openpyxl.
    """
    empresa_id = request.GET.get('empresa', '').strip()
    faena = request.GET.get('faena', '').strip()
    tipo_servicio = request.GET.get('tipo_servicio', '').strip()
    tipo_residuo = request.GET.get('tipo_residuo', '').strip()
    desde = request.GET.get('desde', '').strip() or request.GET.get('fecha_desde', '').strip()
    hasta = request.GET.get('hasta', '').strip() or request.GET.get('fecha_hasta', '').strip()

    qs = TicketRetiro.objects.select_related('empresa').prefetch_related('detalles__material').all().order_by('fecha', 'id')

    empresa_nombre = "TODOS"
    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        qs = qs.filter(empresa=request.user.empresa)
        empresa_nombre = request.user.empresa.nombre
    elif empresa_id:
        qs = qs.filter(empresa_id=empresa_id)
        emp = Empresa.objects.filter(id=empresa_id).first()
        if emp:
            empresa_nombre = emp.nombre

    if faena:
        qs = qs.filter(faena_area__icontains=faena)
    if tipo_servicio:
        qs = qs.filter(tipo_servicio__icontains=tipo_servicio)
    if tipo_residuo:
        qs = qs.filter(tipo_residuo__icontains=tipo_residuo)
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)

    tickets = list(qs)

    materiales_presentes_ids = DetalleMaterialTicket.objects.filter(ticket__in=qs).values_list('material_id', flat=True).distinct()
    columnas_materiales = list(CatalogoMaterialRetiro.objects.filter(id__in=materiales_presentes_ids).order_by('orden', 'nombre'))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado Retiros"

    # Estilos Corporativos Redimir (#006BB8 y #95BF3C)
    font_title = Font(name="Calibri", size=14, bold=True, color="006BB8")
    font_sub = Font(name="Calibri", size=10, italic=True, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")
    font_data = Font(name="Calibri", size=10)

    fill_header_main = PatternFill(start_color="006BB8", end_color="006BB8", fill_type="solid")
    fill_header_mat = PatternFill(start_color="95BF3C", end_color="95BF3C", fill_type="solid")
    fill_total = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    total_border = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000')
    )

    ws.merge_cells('A1:G1')
    ws['A1'] = "REDIMIR SPA — INFORME CONSOLIDADO DE RETIRO DE RESIDUOS"
    ws['A1'].font = font_title

    subtitulo_info = f"Empresa / Cliente: {empresa_nombre} | Faena: {faena or 'TODAS'} | Período: {desde or 'Inicio'} al {hasta or 'Cierre'}"
    ws.merge_cells('A2:G2')
    ws['A2'] = subtitulo_info
    ws['A2'].font = font_sub

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[4].height = 26

    headers = ['N° FOLIO / TICKET', 'FECHA', 'EMPRESA / CLIENTE', 'TIPO SERVICIO', 'TIPO RESIDUO', 'PESO TOTAL BÁSCULA (kg)']
    
    col_mapping = []
    for mat in columnas_materiales:
        if mat.unidad_medida == 'kg_un':
            headers.append(f"{mat.nombre} (kg)")
            col_mapping.append((mat, 'peso'))
            headers.append(f"Cant. {mat.nombre} (un)")
            col_mapping.append((mat, 'cantidad'))
        elif mat.unidad_medida == 'un':
            headers.append(f"{mat.nombre} (un)")
            col_mapping.append((mat, 'cantidad'))
        else:
            headers.append(f"{mat.nombre} (kg)")
            col_mapping.append((mat, 'peso'))

    headers.append('OBSERVACIONES')

    start_row = 4
    for c_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(start_row, c_idx, h_text)
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if c_idx <= 6:
            cell.fill = fill_header_main
        else:
            cell.fill = fill_header_mat
        cell.border = thin_border

    curr_row = start_row + 1
    for t in tickets:
        det_map = {d.material_id: d for d in t.detalles.all()}
        
        ws.cell(curr_row, 1, t.numero_ticket).alignment = Alignment(horizontal="center")
        ws.cell(curr_row, 2, t.fecha.strftime('%d/%m/%Y')).alignment = Alignment(horizontal="center")
        ws.cell(curr_row, 3, t.empresa.nombre)
        ws.cell(curr_row, 4, t.tipo_servicio or '')
        ws.cell(curr_row, 5, t.tipo_residuo or '')
        
        c_peso = ws.cell(curr_row, 6, float(t.peso_total_ticket or 0))
        c_peso.number_format = '#,##0.0'
        c_peso.alignment = Alignment(horizontal="right")

        c_idx = 7
        for mat, metric in col_mapping:
            det = det_map.get(mat.id)
            val = None
            if det:
                if metric == 'peso' and det.peso_kg and det.peso_kg > 0:
                    val = float(det.peso_kg)
                elif metric == 'cantidad' and det.cantidad_unidades and det.cantidad_unidades > 0:
                    val = float(det.cantidad_unidades)
            
            cell_m = ws.cell(curr_row, c_idx, val if val is not None else "")
            if val is not None:
                cell_m.number_format = '#,##0.0' if metric == 'peso' else '#,##0'
                cell_m.alignment = Alignment(horizontal="right")
            c_idx += 1

        ws.cell(curr_row, c_idx, t.observaciones or "")

        for c in range(1, len(headers) + 1):
            cell = ws.cell(curr_row, c)
            cell.font = font_data
            cell.border = thin_border

        curr_row += 1

    total_row = curr_row
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    cell_tot_lbl = ws.cell(total_row, 1, "TOTALES ACUMULADOS")
    cell_tot_lbl.font = font_total
    cell_tot_lbl.alignment = Alignment(horizontal="center", vertical="center")
    cell_tot_lbl.fill = fill_total

    col_letter_bascula = get_column_letter(6)
    cell_tot_bascula = ws.cell(total_row, 6, f"=SUM({col_letter_bascula}{start_row + 1}:{col_letter_bascula}{total_row - 1})")
    cell_tot_bascula.font = font_total
    cell_tot_bascula.fill = fill_total
    cell_tot_bascula.number_format = '#,##0.0'
    cell_tot_bascula.border = total_border

    for c_idx in range(7, len(headers)):
        c_letter = get_column_letter(c_idx)
        cell_tot = ws.cell(total_row, c_idx, f"=SUM({c_letter}{start_row + 1}:{c_letter}{total_row - 1})")
        cell_tot.font = font_total
        cell_tot.fill = fill_total
        cell_tot.number_format = '#,##0.0'
        cell_tot.border = total_border

    cell_empty = ws.cell(total_row, len(headers), "")
    cell_empty.fill = fill_total
    cell_empty.border = total_border

    for c in range(1, 6):
        ws.cell(total_row, c).border = total_border

    ws.row_dimensions[total_row].height = 24

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row in [1, 2]:
                continue
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    filename_empresa = re.sub(r'[^a-zA-Z0-9_-]', '_', empresa_nombre)
    timestamp = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    filename = f"REDIMIR_Informe_Retiros_{filename_empresa}_{timestamp}.xlsx"

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def api_materiales_retiro(request):
    materiales = CatalogoMaterialRetiro.objects.filter(activo=True).order_by('orden', 'nombre').values(
        'id', 'nombre', 'codigo', 'categoria', 'unidad_medida'
    )
    return JsonResponse({'materiales': list(materiales)})


@login_required
def galeria_fotos(request):
    """
    Galería consolidada de fotos y comprobantes con filtros por fecha, categoría y empresa.
    """
    empresa_id = request.GET.get('empresa', '').strip()
    categoria = request.GET.get('categoria', '').strip()
    desde = request.GET.get('desde', '').strip() or request.GET.get('fecha_desde', '').strip()
    hasta = request.GET.get('hasta', '').strip() or request.GET.get('fecha_hasta', '').strip()
    q = request.GET.get('q', '').strip()

    # Query base de tickets con respaldo
    tickets_qs = TicketRetiro.objects.select_related('empresa').exclude(respaldo_ticket='').exclude(respaldo_ticket__isnull=True).order_by('-fecha', '-id')

    if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
        tickets_qs = tickets_qs.filter(empresa=request.user.empresa)
        empresas_list = [request.user.empresa]
    else:
        empresas_list = Empresa.objects.filter(activa=True).order_by('nombre')
        if empresa_id:
            tickets_qs = tickets_qs.filter(empresa_id=empresa_id)

    if desde:
        tickets_qs = tickets_qs.filter(fecha__gte=desde)
    if hasta:
        tickets_qs = tickets_qs.filter(fecha__lte=hasta)
    if q:
        tickets_qs = tickets_qs.filter(
            Q(numero_ticket__icontains=q) |
            Q(faena_area__icontains=q) |
            Q(tipo_servicio__icontains=q) |
            Q(observaciones__icontains=q)
        )

    fotos = []
    categorias_encontradas = set()

    for t in tickets_qs:
        es_img = t.es_imagen_respaldo
        cat = t.tipo_servicio or 'Comprobante Báscula'
        categorias_encontradas.add(cat)

        if categoria and cat != categoria:
            continue

        fotos.append({
            'id': t.id,
            'url': t.respaldo_ticket.url if t.respaldo_ticket else '',
            'es_imagen': es_img,
            'numero_ticket': t.numero_ticket,
            'empresa': t.empresa.nombre,
            'empresa_id': t.empresa_id,
            'fecha': t.fecha,
            'faena': t.faena_area or 'Principal',
            'categoria': cat,
            'tipo_residuo': t.tipo_residuo,
            'peso_total': t.peso_total_ticket,
            'observaciones': t.observaciones,
            'tipo_fuente': 'Ticket Báscula',
            'detalle_url': f"/retiros/{t.id}/",
        })

    # Evidencias fotográficas de módulos si existen
    try:
        # Fotos RSD
        fotos_rsd_qs = FotoRegistroRSD.objects.select_related('registro__servicio__empresa').order_by('-fecha_subida')
        if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
            fotos_rsd_qs = fotos_rsd_qs.filter(registro__servicio__empresa=request.user.empresa)
        elif empresa_id:
            fotos_rsd_qs = fotos_rsd_qs.filter(registro__servicio__empresa_id=empresa_id)
        if desde:
            fotos_rsd_qs = fotos_rsd_qs.filter(registro__servicio__fecha_solicitud__date__gte=desde)
        if hasta:
            fotos_rsd_qs = fotos_rsd_qs.filter(registro__servicio__fecha_solicitud__date__lte=hasta)

        cat_rsd = 'Residuos Domiciliarios (RSD)'
        categorias_encontradas.add(cat_rsd)
        if not categoria or categoria == cat_rsd:
            for f in fotos_rsd_qs[:50]:
                srv = f.registro.servicio
                fotos.append({
                    'id': f"rsd_{f.id}",
                    'url': f.foto.url,
                    'es_imagen': True,
                    'numero_ticket': f"SRV #{srv.pk}",
                    'empresa': srv.empresa.nombre,
                    'empresa_id': srv.empresa_id,
                    'fecha': srv.fecha_solicitud.date(),
                    'faena': srv.direccion,
                    'categoria': cat_rsd,
                    'tipo_residuo': 'RSD / Basura',
                    'peso_total': f.registro.cantidad_kg,
                    'observaciones': f.registro.observaciones,
                    'tipo_fuente': 'Registro RSD',
                    'detalle_url': f"/servicios/{srv.pk}/",
                })
    except Exception:
        pass

    try:
        # Fotos Reciclables
        fotos_rec_qs = FotoRegistroReciclables.objects.select_related('registro__servicio__empresa').order_by('-fecha_subida')
        if getattr(request.user, 'rol', '') == 'empresa' and getattr(request.user, 'empresa', None):
            fotos_rec_qs = fotos_rec_qs.filter(registro__servicio__empresa=request.user.empresa)
        elif empresa_id:
            fotos_rec_qs = fotos_rec_qs.filter(registro__servicio__empresa_id=empresa_id)
        if desde:
            fotos_rec_qs = fotos_rec_qs.filter(registro__servicio__fecha_solicitud__date__gte=desde)
        if hasta:
            fotos_rec_qs = fotos_rec_qs.filter(registro__servicio__fecha_solicitud__date__lte=hasta)

        cat_rec = 'Puntos Verdes / Reciclables'
        categorias_encontradas.add(cat_rec)
        if not categoria or categoria == cat_rec:
            for f in fotos_rec_qs[:50]:
                srv = f.registro.servicio
                fotos.append({
                    'id': f"rec_{f.id}",
                    'url': f.foto.url,
                    'es_imagen': True,
                    'numero_ticket': f"SRV #{srv.pk}",
                    'empresa': srv.empresa.nombre,
                    'empresa_id': srv.empresa_id,
                    'fecha': srv.fecha_solicitud.date(),
                    'faena': srv.direccion,
                    'categoria': cat_rec,
                    'tipo_residuo': f.registro.get_material_display() if hasattr(f.registro, 'get_material_display') else 'Reciclable',
                    'peso_total': f.registro.cantidad_kg,
                    'observaciones': f.registro.observaciones,
                    'tipo_fuente': 'Registro Reciclaje',
                    'detalle_url': f"/servicios/{srv.pk}/",
                })
    except Exception:
        pass

    # Ordenar fotos cronológicamente descendente
    fotos.sort(key=lambda x: str(x['fecha']), reverse=True)

    paginator = Paginator(fotos, 32)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'fotos': page_obj,
        'total_fotos': len(fotos),
        'empresas': empresas_list,
        'categorias': sorted(list(categorias_encontradas)),
        'filtros': {
            'empresa': empresa_id,
            'categoria': categoria,
            'desde': desde,
            'hasta': hasta,
            'q': q,
        }
    }
    return render(request, 'servicios/galeria.html', context)
